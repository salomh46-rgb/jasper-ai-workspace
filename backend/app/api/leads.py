import io
import httpx
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
from app.core.database import get_db
from app.api.deps import get_current_user
from app.models.user import User
from app.models.agent import Agent
from app.models.lead import Lead
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/leads", tags=["Leads"])

class LeadStatusUpdate(BaseModel):
    status: str

class SheetsSyncRequest(BaseModel):
    webhook_url: Optional[str] = None
    sheet_id: Optional[str] = "1xmeMSCZmyoheJ9h7M-krzYo5OJkkCkpk71_LluHwb60"

@router.get("/export/excel")
async def export_leads_excel(
    agent_id: Optional[int] = None,
    status: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    stmt = select(Lead).join(Agent).where(Agent.user_id == current_user.id).options(selectinload(Lead.agent))
    if agent_id:
        stmt = stmt.where(Lead.agent_id == agent_id)
    if status:
        stmt = stmt.where(Lead.status == status)
    stmt = stmt.order_by(Lead.id.desc())

    res = await db.execute(stmt)
    leads = res.scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lidlar CRM"

    # Header styling
    headers = ["ID", "Agent (Bot)", "Mijoz Ismi", "Telefon Raqami", "Murojaat / Maqsad", "AI Xulosa", "Holati", "Yaratilgan Vaqt"]
    ws.append(headers)

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    status_labels = {
        "new": "Yangi",
        "in_progress": "Jarayonda",
        "completed": "Bajarildi",
        "canceled": "Bekor qilingan"
    }

    row_font = Font(name="Segoe UI", size=10)
    for idx, l in enumerate(leads, start=2):
        created_str = l.created_at.strftime("%Y-%m-%d %H:%M") if l.created_at else ""
        row_data = [
            l.id,
            l.agent.name if l.agent else "",
            l.customer_name or "Noma'lum",
            l.customer_phone or "",
            l.intent or "",
            l.summary or "",
            status_labels.get(l.status, l.status),
            created_str
        ]
        ws.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=idx, column=col_idx)
            c.font = row_font
            c.border = thin_border
            if col_idx in [1, 4, 7, 8]:
                c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[idx].height = 22

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = "Jasper_CRM_Lidlar.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.post("/sync-sheets")
async def sync_leads_to_google_sheets(
    payload: SheetsSyncRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    stmt = select(Lead).join(Agent).where(Agent.user_id == current_user.id).options(selectinload(Lead.agent)).order_by(Lead.id.desc()).limit(100)
    res = await db.execute(stmt)
    leads = res.scalars().all()

    leads_data = []
    for l in leads:
        leads_data.append({
            "id": l.id,
            "agent_name": l.agent.name if l.agent else "",
            "customer_name": l.customer_name or "Mijoz",
            "customer_phone": l.customer_phone or "",
            "intent": l.intent or "",
            "summary": l.summary or "",
            "status": l.status or "new",
            "created_at": str(l.created_at)
        })

    if payload.webhook_url:
        try:
            async with httpx.AsyncClient(timeout=15.0) as client:
                resp = await client.post(payload.webhook_url, json={"action": "sync_leads", "sheet_id": payload.sheet_id, "leads": leads_data})
                return {"status": "success", "message": f"{len(leads_data)} ta lid Google Sheets ga yuborildi!", "status_code": resp.status_code}
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Google Sheets ga yuborishda xatolik: {str(e)}")

    return {
        "status": "success",
        "sheet_id": payload.sheet_id or "1xmeMSCZmyoheJ9h7M-krzYo5OJkkCkpk71_LluHwb60",
        "total_leads_prepared": len(leads_data),
        "message": f"Google Sheets ({payload.sheet_id}) uchun {len(leads_data)} ta lid tayyorlandi!"
    }

@router.get("")
async def list_leads(
    agent_id: Optional[int] = None,
    status: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    stmt = select(Lead).join(Agent).where(Agent.user_id == current_user.id).options(selectinload(Lead.agent))
    if agent_id:
        stmt = stmt.where(Lead.agent_id == agent_id)
    if status:
        stmt = stmt.where(Lead.status == status)

    stmt = stmt.order_by(Lead.id.desc())
    res = await db.execute(stmt)
    leads = res.scalars().all()

    result = []
    for l in leads:
        result.append({
            "id": l.id,
            "agent_id": l.agent_id,
            "agent_name": l.agent.name if l.agent else "Noma'lum",
            "customer_tg_id": l.customer_tg_id,
            "customer_name": l.customer_name,
            "customer_phone": l.customer_phone,
            "intent": l.intent,
            "summary": l.summary,
            "status": l.status,
            "created_at": l.created_at
        })
    return result

@router.put("/{lead_id}")
async def update_lead_status(lead_id: int, payload: LeadStatusUpdate, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    stmt = select(Lead).join(Agent).where(Lead.id == lead_id, Agent.user_id == current_user.id)
    res = await db.execute(stmt)
    lead = res.scalar_one_or_none()
    if not lead:
        raise HTTPException(status_code=404, detail="Lid topilmadi")

    lead.status = payload.status
    await db.commit()
    return {"status": "success", "message": "Lid holati yangilandi"}
